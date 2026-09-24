"""Text and image extraction from uploaded reference documents (PDF, Word).

Image extraction pulls out embedded photos (e.g. label images attached to a
COLA application PDF/docx) so they can be run through the same OCR pipeline
as a manually uploaded label photo, without the user having to re-upload
something that's already in the document.
"""
import io
import os

import pypdf
from docx import Document
from PIL import Image

ALLOWED_EXTENSIONS = {".pdf", ".docx"}

# Filters out small embedded images (logos, bullet icons, letterhead
# graphics) that aren't a real label photo.
MIN_EMBEDDED_IMAGE_DIMENSION = 150

_LOCKED_PDF_ERROR = ValueError(
    "This PDF is password-protected and can't be opened automatically. "
    "Please upload an unprotected PDF, or remove the password first."
)


def _readable_pdf_reader(file_storage) -> pypdf.PdfReader:
    """Open ``file_storage`` as a PdfReader, transparently unlocking it if
    it's encrypted with an empty password. Raises ValueError if it's
    genuinely password-protected."""
    file_storage.stream.seek(0)
    reader = pypdf.PdfReader(file_storage.stream)

    if reader.is_encrypted:
        try:
            result = reader.decrypt("")
        except Exception as e:
            raise _LOCKED_PDF_ERROR from e
        if not result:
            raise _LOCKED_PDF_ERROR

    return reader


def _extract_pdf_text(file_storage) -> str:
    reader = _readable_pdf_reader(file_storage)

    try:
        pages = [page.extract_text() or "" for page in reader.pages]
    except pypdf.errors.FileNotDecryptedError as e:
        raise _LOCKED_PDF_ERROR from e
    except pypdf.errors.DependencyError as e:
        raise ValueError(
            "This PDF uses encryption that couldn't be read (missing the 'cryptography' package "
            "on the server). Try re-saving the PDF without security/encryption, or contact the app "
            "maintainer to install the missing dependency."
        ) from e

    return "\n".join(pages)


def _extract_docx_text(file_storage) -> str:
    file_storage.stream.seek(0)
    document = Document(file_storage.stream)
    return "\n".join(p.text for p in document.paragraphs)


def extract_text(file_storage) -> str:
    """Extract plain text from an uploaded .pdf or .docx file.

    Raises ValueError for unsupported file types.
    """
    filename = file_storage.filename or ""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return _extract_pdf_text(file_storage)
    if ext == ".docx":
        return _extract_docx_text(file_storage)
    raise ValueError(f"Unsupported document type '{ext or filename}'. Please upload a PDF or Word (.docx) file.")


def _large_enough(image: Image.Image) -> bool:
    return min(image.width, image.height) >= MIN_EMBEDDED_IMAGE_DIMENSION


def _images_from_pdf(file_storage) -> list[Image.Image]:
    try:
        reader = _readable_pdf_reader(file_storage)
    except ValueError:
        return []

    images = []
    for page in reader.pages:
        try:
            page_images = list(page.images)
        except Exception:
            continue
        for image_file in page_images:
            try:
                pil_image = image_file.image
            except Exception:
                continue
            if pil_image is None or not _large_enough(pil_image):
                continue
            images.append(pil_image.convert("RGB"))
    return images


def _images_from_docx(file_storage) -> list[Image.Image]:
    file_storage.stream.seek(0)
    document = Document(file_storage.stream)

    images = []
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        try:
            pil_image = Image.open(io.BytesIO(rel.target_part.blob))
            pil_image.load()
        except Exception:
            continue
        if not _large_enough(pil_image):
            continue
        images.append(pil_image.convert("RGB"))
    return images


def extract_images(file_storage) -> list[Image.Image]:
    """Extract embedded photos from a .pdf or .docx (e.g. label images
    attached to a COLA application). Best-effort: an unsupported file type,
    a genuinely password-locked PDF, or any per-image read failure just
    yields fewer/no images rather than raising, since this is a convenience
    on top of the required text extraction, not something that should block
    the whole upload.
    """
    filename = file_storage.filename or ""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return _images_from_pdf(file_storage)
    if ext == ".docx":
        return _images_from_docx(file_storage)
    return []
