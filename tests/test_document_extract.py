import io

from docx import Document
from PIL import Image
from werkzeug.datastructures import FileStorage

from app.documents.extract import extract_images, extract_text


def _pdf_with_image(size=(300, 300)) -> FileStorage:
    img = Image.new("RGB", size, "white")
    buf = io.BytesIO()
    img.save(buf, "PDF", resolution=150.0)
    buf.seek(0)
    return FileStorage(stream=buf, filename="test.pdf", content_type="application/pdf")


def _docx_with_image(size=(300, 300)) -> FileStorage:
    img = Image.new("RGB", size, "white")
    img_buf = io.BytesIO()
    img.save(img_buf, "PNG")
    img_buf.seek(0)

    doc = Document()
    doc.add_paragraph("Brand Name: Test Brand")
    doc.add_picture(img_buf)
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return FileStorage(
        stream=out,
        filename="test.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def test_extract_images_from_pdf_returns_embedded_image():
    images = extract_images(_pdf_with_image())
    assert len(images) == 1
    assert images[0].width >= 150 and images[0].height >= 150


def test_extract_images_from_docx_returns_embedded_image():
    images = extract_images(_docx_with_image())
    assert len(images) == 1


def test_extract_images_filters_small_images():
    # Small enough to be a logo/icon rather than an actual label photo.
    assert extract_images(_pdf_with_image(size=(50, 50))) == []


def test_extract_images_unsupported_type_returns_empty_list():
    fs = FileStorage(stream=io.BytesIO(b"not a document"), filename="notes.txt", content_type="text/plain")
    assert extract_images(fs) == []


def test_extract_text_and_extract_images_can_both_read_same_file():
    # Both need to read the same underlying stream -- confirm one doesn't
    # leave it exhausted/unreadable for the other.
    fs = _pdf_with_image()
    extract_text(fs)
    images = extract_images(fs)
    assert len(images) == 1
