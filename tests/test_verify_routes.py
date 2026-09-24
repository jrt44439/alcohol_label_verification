import io

from docx import Document

from app.models import product as product_model
from app.models import verification as verification_model


def _cola_docx(brand="Test Brand", class_type="Bourbon Whiskey") -> io.BytesIO:
    doc = Document()
    doc.add_paragraph(f"Brand Name: {brand}")
    doc.add_paragraph(f"Class and Type: {class_type}")
    doc.add_paragraph("Name and Address of Producer: Test Distillery, Frankfort, KY")
    doc.add_paragraph("Country of Origin: United States")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def test_save_product_create(client):
    response = client.post(
        "/verify/save-product",
        data={
            "action": "create",
            "name": "New From Verify",
            "brand_name": "Summit Peak",
            "class_type": "Rye Whiskey",
            "alcohol_content": "47% ALC/VOL",
            "net_contents": "750 mL",
            "producer_info": "Summit Peak Distillers, Denver, CO",
            "country_of_origin": "",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"New From Verify" in response.data

    with client.application.app_context():
        products = product_model.get_all()
    assert any(p.name == "New From Verify" and p.brand_name == "Summit Peak" for p in products)


def test_save_product_create_ajax_returns_json_without_redirect(client):
    # The "Add to Product Library" button on a saved COLA result submits
    # with ajax=1 so the page it's on doesn't navigate away.
    response = client.post(
        "/verify/save-product",
        data={
            "action": "create",
            "name": "Old Ridge",
            "brand_name": "Old Ridge",
            "class_type": "Bourbon Whiskey",
            "producer_info": "Test Distillery, Frankfort, KY",
            "ajax": "1",
        },
    )
    assert response.status_code == 200
    assert response.content_type.startswith("application/json")
    payload = response.get_json()
    assert payload["success"] is True
    assert "Old Ridge" in payload["message"]

    with client.application.app_context():
        products = product_model.get_all()
    assert any(p.name == "Old Ridge" and p.producer_info == "Test Distillery, Frankfort, KY" for p in products)


def test_save_product_create_ajax_missing_name_returns_json_error(client):
    response = client.post(
        "/verify/save-product",
        data={"action": "create", "brand_name": "Old Ridge", "ajax": "1"},
    )
    assert response.status_code == 200
    payload = response.get_json()
    assert payload["success"] is False
    assert "name" in payload["message"].lower()


def test_save_product_create_requires_name(client):
    response = client.post(
        "/verify/save-product",
        data={"action": "create", "brand_name": "Summit Peak"},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Enter a name for the new product" in response.data

    with client.application.app_context():
        products = product_model.get_all()
    assert products == []


def test_save_product_update(client):
    with client.application.app_context():
        product_id = product_model.create(
            {
                "name": "Existing Product",
                "brand_name": "Old Brand",
                "class_type": "Vodka",
                "alcohol_content": "40% ALC/VOL",
                "net_contents": "1 L",
                "producer_info": "Old Producer",
                "country_of_origin": None,
            }
        )

    response = client.post(
        "/verify/save-product",
        data={
            "action": "update",
            "product_id": str(product_id),
            "brand_name": "Corrected Brand",
            "class_type": "Vodka",
            "alcohol_content": "40% ALC/VOL",
            "net_contents": "1 L",
            "producer_info": "Corrected Producer",
            "country_of_origin": "",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Updated" in response.data

    with client.application.app_context():
        updated = product_model.get_by_id(product_id)
    assert updated.name == "Existing Product"
    assert updated.brand_name == "Corrected Brand"
    assert updated.producer_info == "Corrected Producer"


def test_save_product_update_missing_product(client):
    response = client.post(
        "/verify/save-product",
        data={"action": "update", "product_id": "999", "brand_name": "X"},
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Product not found" in response.data


def test_compare_route_applies_manual_field_override(client):
    form_data = {
        "item_count": "1",
        "reference_mode_0": "adhoc",
        "raw_text_0": "OLD RIDGE",
        "extracted_brand_name_0": "Old Ridge",
        "expected_brand_name_0": "Old Ridge",
        "extracted_class_type_0": "Bourbon",
        "expected_class_type_0": "Bourbon",
        "extracted_alcohol_content_0": "45% ALC/VOL",
        "expected_alcohol_content_0": "45% ALC/VOL",
        "extracted_net_contents_0": "750 mL",
        "expected_net_contents_0": "750 mL",
        "extracted_producer_info_0": "Old Ridge Distillery",
        "expected_producer_info_0": "Old Ridge Distillery",
        "extracted_country_of_origin_0": "",
        "expected_country_of_origin_0": "",
        "extracted_health_warning_text_0": "",
        "expected_health_warning_text_0": "",
        # class_type would naturally compute as MATCH -- force it to MISMATCH
        # without touching either value.
        "override_status_0_class_type": "mismatch",
    }
    response = client.post("/verify/compare", data=form_data, follow_redirects=True)
    assert response.status_code == 200
    assert b"(Manual)" in response.data
    assert b'value="Old Ridge"' in response.data
    assert b'value="Bourbon"' in response.data


def test_run_cola_mode_without_document_flashes_error(client):
    # COLA validation happens before any OCR, so a dummy image is fine here.
    data = {
        "batch_reference_mode": "cola",
        "label_images": (io.BytesIO(b"fake image bytes"), "label.png"),
    }
    response = client.post("/verify/run", data=data, content_type="multipart/form-data", follow_redirects=True)
    assert response.status_code == 200
    assert b"Please choose a COLA application" in response.data


def test_run_cola_mode_with_unsupported_file_type_flashes_error(client):
    data = {
        "batch_reference_mode": "cola",
        "label_images": (io.BytesIO(b"fake image bytes"), "label.png"),
        "cola_document": (io.BytesIO(b"not a real document"), "application.txt"),
    }
    response = client.post("/verify/run", data=data, content_type="multipart/form-data", follow_redirects=True)
    assert response.status_code == 200
    assert b"Unsupported document type" in response.data


def test_run_cola_batch_without_documents_flashes_error(client):
    response = client.post("/verify/run-cola-batch", data={}, follow_redirects=True)
    assert response.status_code == 200
    assert b"Please choose at least one COLA application" in response.data
    with client.application.app_context():
        assert verification_model.get_all() == []


def test_run_cola_batch_saves_valid_documents_and_warns_on_bad_ones(client):
    data = {
        "cola_documents": [
            (io.BytesIO(b"not a real document"), "bad.txt"),
            (_cola_docx(), "good.docx"),
        ],
    }
    response = client.post(
        "/verify/run-cola-batch", data=data, content_type="multipart/form-data", follow_redirects=True
    )
    assert response.status_code == 200
    assert b"bad.txt" in response.data
    assert b"Saved 1 verification result" in response.data

    with client.application.app_context():
        saved = verification_model.get_all()
    assert len(saved) == 1
    assert saved[0].source_filename == "good.docx"


def test_run_cola_batch_skips_duplicate_brand_name_in_same_batch(client):
    data = {
        "cola_documents": [
            (_cola_docx(brand="Old Ridge"), "first.docx"),
            (_cola_docx(brand="Old Ridge"), "second.docx"),
        ],
    }
    response = client.post(
        "/verify/run-cola-batch", data=data, content_type="multipart/form-data", follow_redirects=True
    )
    assert response.status_code == 200
    assert b"Saved 1 verification result" in response.data
    assert b"Skipped 1 duplicate application" in response.data
    assert b"second.docx" in response.data

    with client.application.app_context():
        saved = verification_model.get_all()
    assert len(saved) == 1
    assert saved[0].source_filename == "first.docx"


def test_run_cola_batch_skips_duplicate_against_previously_saved(client):
    client.post(
        "/verify/run-cola-batch",
        data={"cola_documents": [(_cola_docx(brand="Old Ridge"), "first.docx")]},
        content_type="multipart/form-data",
    )

    response = client.post(
        "/verify/run-cola-batch",
        data={"cola_documents": [(_cola_docx(brand="Old Ridge"), "second.docx")]},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Skipped 1 duplicate application" in response.data

    with client.application.app_context():
        saved = verification_model.get_all()
    assert len(saved) == 1


def test_run_cola_batch_allows_different_brand_names(client):
    data = {
        "cola_documents": [
            (_cola_docx(brand="Old Ridge"), "first.docx"),
            (_cola_docx(brand="New Summit"), "second.docx"),
        ],
    }
    response = client.post(
        "/verify/run-cola-batch", data=data, content_type="multipart/form-data", follow_redirects=True
    )
    assert response.status_code == 200
    assert b"Saved 2 verification result" in response.data

    with client.application.app_context():
        saved = verification_model.get_all()
    assert len(saved) == 2


def test_view_saved_result_renders_stored_data(client):
    data = {"cola_documents": [(_cola_docx(brand="Old Ridge", class_type="Bourbon"), "old_ridge.docx")]}
    client.post("/verify/run-cola-batch", data=data, content_type="multipart/form-data")

    with client.application.app_context():
        verification_id = verification_model.get_all()[0].id

    response = client.get(f"/verify/results/{verification_id}")
    assert response.status_code == 200
    assert b"old_ridge.docx" in response.data
    # No embedded photos in this document, so every field is unverifiable.
    assert b"NOT FOUND" in response.data
    assert b"no embedded label photos" in response.data


def test_view_saved_result_exposes_application_fields_for_product_library_save(client):
    data = {"cola_documents": [(_cola_docx(brand="Old Ridge", class_type="Bourbon Whiskey"), "old_ridge.docx")]}
    client.post("/verify/run-cola-batch", data=data, content_type="multipart/form-data")

    with client.application.app_context():
        verification_id = verification_model.get_all()[0].id

    response = client.get(f"/verify/results/{verification_id}")
    assert response.status_code == 200
    assert b"Add to Product Library" in response.data
    assert b'"brand_name": "Old Ridge"' in response.data
    assert b'"class_type": "Bourbon Whiskey"' in response.data
    assert b"Test Distillery" in response.data


def test_save_product_create_from_cola_application_fields(client):
    # Simulates the "Add to Product Library" button's client-side form
    # submission, using the application's own extracted values.
    response = client.post(
        "/verify/save-product",
        data={
            "action": "create",
            "name": "Old Ridge",
            "brand_name": "Old Ridge",
            "class_type": "Bourbon Whiskey",
            "alcohol_content": "",
            "net_contents": "",
            "producer_info": "Test Distillery, Frankfort, KY",
            "country_of_origin": "United States",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Saved &#39;Old Ridge&#39; to the product library" in response.data or b"Old Ridge" in response.data

    with client.application.app_context():
        products = product_model.get_all()
    assert any(
        p.name == "Old Ridge" and p.producer_info == "Test Distillery, Frankfort, KY" for p in products
    )


def test_view_saved_result_missing_redirects_to_list(client):
    response = client.get("/verify/results/999", follow_redirects=True)
    assert response.status_code == 200
    assert b"Saved verification not found" in response.data


def test_delete_saved_result(client):
    data = {"cola_documents": [(_cola_docx(), "to_delete.docx")]}
    client.post("/verify/run-cola-batch", data=data, content_type="multipart/form-data")
    with client.application.app_context():
        verification_id = verification_model.get_all()[0].id

    response = client.post(f"/verify/results/{verification_id}/delete", follow_redirects=True)
    assert response.status_code == 200
    with client.application.app_context():
        assert verification_model.get_by_id(verification_id) is None


def test_update_saved_result_persists_correction(client):
    photo = {
        "index": 0,
        "raw_text": "OLD RIDGE\nBourbon\n45% ALC/VOL\n750 mL\nOld Ridge Distillery",
        "extracted": {
            "brand_name": None,  # misread as blank
            "class_type": "Bourbon",
            "alcohol_content": "45% ALC/VOL",
            "net_contents": "750 mL",
            "producer_info": "Old Ridge Distillery",
            "country_of_origin": None,
            "health_warning_text": "GOVERNMENT WARNING: ...",
        },
        "expected": {
            "brand_name": "Old Ridge",
            "class_type": "Bourbon",
            "alcohol_content": "45% ALC/VOL",
            "net_contents": "750 mL",
            "producer_info": "Old Ridge Distillery",
            "country_of_origin": None,
            "health_warning_text": "GOVERNMENT WARNING: ...",
        },
        "comparisons_by_field": {},
        "warning_caps_ok": True,
        "warning_bold_ok": True,
    }
    field_summary = {
        "brand_name": {"status": "not_found", "expected": "Old Ridge"},
        "class_type": {"status": "match", "value": "Bourbon", "source_index": 0, "expected": "Bourbon"},
        "alcohol_content": {"status": "match", "value": "45% ALC/VOL", "source_index": 0, "expected": "45% ALC/VOL"},
        "net_contents": {"status": "match", "value": "750 mL", "source_index": 0, "expected": "750 mL"},
        "producer_info": {
            "status": "match",
            "value": "Old Ridge Distillery",
            "source_index": 0,
            "expected": "Old Ridge Distillery",
        },
        "country_of_origin": {"status": "match", "value": None, "source_index": 0, "expected": None},
        "health_warning_text": {
            "status": "match",
            "value": "GOVERNMENT WARNING: ...",
            "source_index": 0,
            "expected": "GOVERNMENT WARNING: ...",
        },
    }

    with client.application.app_context():
        vid = verification_model.create("test.docx", field_summary, True, True, [photo])

    form_data = {
        "raw_text_0": photo["raw_text"],
        "warning_caps_ok_0": "true",
        "warning_bold_ok_0": "true",
        "extracted_brand_name_0": "Old Ridge",
        "expected_brand_name_0": "Old Ridge",
        "extracted_class_type_0": "Bourbon",
        "expected_class_type_0": "Bourbon",
        "extracted_alcohol_content_0": "45% ALC/VOL",
        "expected_alcohol_content_0": "45% ALC/VOL",
        "extracted_net_contents_0": "750 mL",
        "expected_net_contents_0": "750 mL",
        "extracted_producer_info_0": "Old Ridge Distillery",
        "expected_producer_info_0": "Old Ridge Distillery",
        "extracted_country_of_origin_0": "",
        "expected_country_of_origin_0": "",
        "extracted_health_warning_text_0": "GOVERNMENT WARNING: ...",
        "expected_health_warning_text_0": "GOVERNMENT WARNING: ...",
    }
    response = client.post(f"/verify/results/{vid}/update", data=form_data, follow_redirects=True)
    assert response.status_code == 200
    assert b"Saved corrections" in response.data

    with client.application.app_context():
        updated = verification_model.get_by_id(vid)
    assert updated.field_summary["brand_name"]["status"] == "match"
    assert updated.overall_status == "pass"


def test_update_saved_result_applies_manual_override(client):
    field_summary = {
        "brand_name": {"status": "match", "value": "Old Ridge", "source_index": 0, "expected": "Old Ridge"},
        "class_type": {"status": "match", "value": "Bourbon", "source_index": 0, "expected": "Bourbon"},
        "alcohol_content": {"status": "match", "value": "45% ALC/VOL", "source_index": 0, "expected": "45% ALC/VOL"},
        "net_contents": {"status": "match", "value": "750 mL", "source_index": 0, "expected": "750 mL"},
        "producer_info": {"status": "match", "value": "Old Ridge Distillery", "source_index": 0, "expected": "Old Ridge Distillery"},
        "country_of_origin": {"status": "match", "value": None, "source_index": 0, "expected": None},
        "health_warning_text": {"status": "match", "value": "GOVERNMENT WARNING: ...", "source_index": 0, "expected": "GOVERNMENT WARNING: ..."},
    }
    with client.application.app_context():
        vid = verification_model.create("test.docx", field_summary, True, True, [])

    response = client.post(
        f"/verify/results/{vid}/update",
        data={"override_status_brand_name": "mismatch"},
        follow_redirects=True,
    )
    assert response.status_code == 200

    with client.application.app_context():
        updated = verification_model.get_by_id(vid)
    # The value itself is untouched -- only the status was overridden.
    assert updated.field_summary["brand_name"]["status"] == "mismatch"
    assert updated.field_summary["brand_name"]["value"] == "Old Ridge"
    assert updated.field_summary["brand_name"]["manual_override"] is True
    assert updated.field_summary["brand_name"]["computed_status"] == "match"
    # Untouched fields keep their computed status and no override flag.
    assert updated.field_summary["class_type"]["status"] == "match"
    assert updated.field_summary["class_type"]["manual_override"] is False
    assert updated.overall_status == "needs_review"


def test_update_saved_result_missing_redirects_to_list(client):
    response = client.post("/verify/results/999/update", data={}, follow_redirects=True)
    assert response.status_code == 200
    assert b"Saved verification not found" in response.data
